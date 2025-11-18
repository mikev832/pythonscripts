import os
import re
import sys
import time
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import load_workbook

from docx import Document  # factory function (used to open the template)
from docx.document import Document as DocxDocument  # actual class for isinstance checks
from docx.text.paragraph import Paragraph
from docx.table import _Cell, Table

import logging
from logging.handlers import RotatingFileHandler

# ---------- Fixed paths ----------
EXCEL_PATH = r"C:\Scripts\azureclivmdumproughdraft.xlsx"
TEMPLATE_PATH = r"C:\Scripts\Template.docx"
OUTPUT_DIR = r"C:\Scripts\VMdocs"
WORKSHEET_NAME: Optional[str] = None  # e.g., "Sheet1" to force a sheet
LOG_FILE = r"C:\Scripts\generate_server_docs.log"

# ---------- Placeholder -> Excel column mapping ----------
PLACEHOLDER_MAP: Dict[str, str] = {
    "<server_name>": "server_name",
    "<subscription_name>": "subscription_name",
    "<VNET/subnet>": "VNET/subnet",     # likely missing; becomes blank
    "<PrivateIP>": "PrivateIP",
    "<PublicIP>": "PublicIP",
    "<region name>": "region_name",
    "<OS>": "OS",
    "<Image>": "Image",
    "<Size>": "Size",
    "<OS_Disk_Name>": "OS_Disk_Name",
    "<OS_Disk_Type>": "OS_Disk_Type",
    "<Resource_Group>": "Resource_Group",
    "<screenshot>": "screenshot",       # likely missing; becomes blank
}

# ---------- Logging ----------
def setup_logging(log_path: str) -> logging.Logger:
    logger = logging.getLogger("docgen")
    logger.setLevel(logging.DEBUG)

    fh = RotatingFileHandler(log_path, maxBytes=1_000_000, backupCount=5, encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(logging.Formatter("%(asctime)s | %(levelname)-8s | %(message)s"))

    ch = logging.StreamHandler(stream=sys.stdout)
    ch.setLevel(logging.INFO)
    ch.setFormatter(logging.Formatter("%(message)s"))

    if not logger.handlers:
        logger.addHandler(fh)
        logger.addHandler(ch)

    logging.getLogger("openpyxl").setLevel(logging.WARNING)
    return logger

logger = setup_logging(LOG_FILE)

# ---------- Utilities ----------
INVALID_FILENAME_CHARS = r'[\\/:*?"<>|]'

def sanitize_filename(name: str, default="UNTITLED") -> str:
    if not name or not str(name).strip():
        return default
    clean = re.sub(INVALID_FILENAME_CHARS, " ", str(name)).strip()
    if not clean or re.fullmatch(r"[\.\s-]+", clean or ""):
        clean = default
    return clean[:150]

def iter_block_items(parent):
    """
    Yield paragraphs and tables in document order.
    Works for Document body, table cells, headers/footers, etc.
    """
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        # header/footer/other containers expose ._element
        parent_elm = parent._element

    for child in parent_elm.iterchildren():
        tag = child.tag
        # Paragraph if tag ends with '}p'
        if tag.endswith('}p'):
            yield Paragraph(child, parent)
        # Table if tag ends with '}tbl'
        elif tag.endswith('}tbl'):
            yield Table(child, parent)

def replace_in_paragraph(paragraph: Paragraph, replacements: Dict[str, str]) -> None:
    """
    Replace placeholders in a paragraph, even when split across runs.
    Strategy: rebuild text across runs; replace; write back as a single run.
    """
    if not paragraph.runs:
        return

    full_text = "".join(run.text for run in paragraph.runs)
    new_text = full_text
    changed = False
    for ph, val in replacements.items():
        if ph in new_text:
            new_text = new_text.replace(ph, val)
            changed = True

    if changed:
        # Remove all runs from the underlying XML, then add one run with the new text
        while paragraph.runs:
            r = paragraph.runs[0]._element
            r.getparent().remove(r)
        paragraph.add_run(new_text)

def replace_in_table(table: Table, replacements: Dict[str, str]) -> None:
    for row in table.rows:
        for cell in row.cells:
            replace_in_container(cell, replacements)

def replace_in_container(container, replacements: Dict[str, str]) -> None:
    """Apply replacements to any container (doc, header, footer, cell)."""
    for block in iter_block_items(container):
        if isinstance(block, Paragraph):
            replace_in_paragraph(block, replacements)
        elif isinstance(block, Table):
            replace_in_table(block, replacements)

def build_replacements(row_dict: Dict[str, str]) -> Dict[str, str]:
    """
    For a given Excel row, build {placeholder: value}; missing/empty => ''.
    Case-insensitive lookup so header case differences don't break mapping.
    """
    ci_map = {k.lower(): k for k in row_dict.keys()}
    rep: Dict[str, str] = {}
    for ph, col in PLACEHOLDER_MAP.items():
        actual_key = ci_map.get(col.lower())
        val = row_dict.get(actual_key, "") if actual_key else ""
        rep[ph] = "" if val is None else str(val)
    return rep

def read_excel_rows(path: str, sheet_name: Optional[str] = None) -> List[Dict[str, str]]:
    """Read Excel and return a list of row dicts keyed by header names."""
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        logger.exception(f"Failed to open Excel workbook: {path}")
        raise

    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        wb.close()
        raise RuntimeError("The worksheet appears to be empty.")

    if not headers:
        wb.close()
        raise RuntimeError("No header row found in the worksheet.")

    headers = [str(h).strip() if h is not None else "" for h in headers]
    data_rows: List[Dict[str, str]] = []
    for row in rows_iter:
        row_dict: Dict[str, str] = {}
        has_data = False
        for idx, cell_val in enumerate(row):
            key = headers[idx] if idx < len(headers) else ""
            if key:
                if cell_val is None:
                    row_dict[key] = ""
                else:
                    text = cell_val if isinstance(cell_val, str) else str(cell_val)
                    row_dict[key] = text
                    if text.strip():
                        has_data = True
        if has_data:
            data_rows.append(row_dict)

    wb.close()
    return data_rows

def ensure_unique_path(base_path: Path) -> Path:
    """If base_path exists, append -1, -2, ... to make it unique."""
    if not base_path.exists():
        return base_path
    stem = base_path.stem
    suffix = base_path.suffix
    parent = base_path.parent
    i = 1
    while True:
        candidate = parent / f"{stem}-{i}{suffix}"
        if not candidate.exists():
            return candidate
        i += 1

# ---------- Simple progress bar ----------
class ProgressBar:
    def __init__(self, total: int, width: int = 40, label: str = "Processing"):
        self.total = max(total, 1)
        self.width = width
        self.label = label
        self.start = time.time()
        self.count = 0
        self._render()

    def _eta(self) -> str:
        elapsed = time.time() - self.start
        if self.count == 0:
            return "--:--"
        rate = elapsed / self.count
        remain = rate * (self.total - self.count)
        m, s = divmod(int(remain), 60)
        return f"{m:02d}:{s:02d}"

    def _render(self):
        pct = self.count / self.total
        filled = int(self.width * pct)
        bar = "#" * filled + "-" * (self.width - filled)
        percent = int(pct * 100)
        msg = f"\r{self.label} [{bar}] {percent:3d}% ({self.count}/{self.total}) ETA {self._eta()}"
        sys.stdout.write(msg)
        sys.stdout.flush()

    def update(self, step: int = 1):
        self.count = min(self.total, self.count + step)
        self._render()
        if self.count == self.total:
            sys.stdout.write("\n")
            sys.stdout.flush()

# ---------- Main generation ----------
def generate_docs():
    # Validate inputs
    if not os.path.exists(EXCEL_PATH):
        msg = f"Excel file not found: {EXCEL_PATH}"
        logger.error(msg)
        raise FileNotFoundError(msg)
    if not os.path.exists(TEMPLATE_PATH):
        msg = f"Template file not found: {TEMPLATE_PATH}"
        logger.error(msg)
        raise FileNotFoundError(msg)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    logger.info("Starting generation...")
    logger.info(f"Excel   : {EXCEL_PATH}")
    logger.info(f"Template: {TEMPLATE_PATH}")
    logger.info(f"Output  : {OUTPUT_DIR}")

    try:
        rows = read_excel_rows(EXCEL_PATH, WORKSHEET_NAME)
    except Exception:
        logger.exception("Failed reading Excel rows.")
        raise

    if not rows:
        logger.warning("No data rows found in Excel. Nothing to do.")
        print("No data rows found in Excel. Nothing to do.")
        return

    total = sum(1 for r in rows if str(r.get("server_name", "")).strip())
    if total == 0:
        logger.warning("No rows with 'server_name' present. Nothing to do.")
        print("No rows with 'server_name' present. Nothing to do.")
        return

    pb = ProgressBar(total, label="Generating docs")
    generated = 0
    skipped = 0
    failures = 0

    for idx, row in enumerate(rows, start=1):
        try:
            server_name = (row.get("server_name") or "").strip()
            if not server_name:
                skipped += 1
                continue

            out_name = sanitize_filename(server_name) + ".docx"
            out_path = ensure_unique_path(Path(OUTPUT_DIR) / out_name)

            # Open template fresh per doc
            try:
                doc = Document(TEMPLATE_PATH)
            except Exception:
                failures += 1
                logger.exception(f"[Row {idx}] Failed to open template for server '{server_name}'.")
                pb.update(1)
                continue

            replacements = build_replacements(row)

            # Replace in body
            try:
                replace_in_container(doc, replacements)
            except Exception:
                failures += 1
                logger.exception(f"[Row {idx}] Replacement failed in body for '{server_name}'.")
                pb.update(1)
                continue

            # Replace in headers/footers
            try:
                for section in doc.sections:
                    if section.header:
                        replace_in_container(section.header, replacements)
                    if section.footer:
                        replace_in_container(section.footer, replacements)
            except Exception:
                failures += 1
                logger.exception(f"[Row {idx}] Replacement failed in header/footer for '{server_name}'.")
                pb.update(1)
                continue

            # Save
            try:
                doc.save(str(out_path))
            except Exception:
                failures += 1
                logger.exception(f"[Row {idx}] Failed to save document '{out_path}'.")
                pb.update(1)
                continue

            generated += 1
            logger.info(f"[OK] {out_path}")
        except Exception:
            failures += 1
            logger.exception(f"[Row {idx}] Unexpected error.")
        finally:
            if (row.get("server_name") or "").strip():
                pb.update(1)

    logger.info(f"Completed. Generated={generated}, Skipped(no server_name)={skipped}, Failures={failures}")
    print(f"Done. Generated {generated} document(s). Skipped {skipped}. Failures {failures}. See log: {LOG_FILE}")

if __name__ == "__main__":
    try:
        generate_docs()
    except Exception as e:
        logger.exception("Fatal error—generation aborted.")
        print(f"Fatal error: {e}\nSee log for details: {LOG_FILE}")
        sys.exit(1)